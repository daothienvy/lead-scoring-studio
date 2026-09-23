import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.utils import get_column_letter

def generate_sales_dashboard():
    src_file = "sample-data/MINDX_Lesson 2_DEMO_synthetic_sales_data_500x20.xlsx"
    out_dir = "outputs/reports"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Sales_Performance_Dashboard_Cleaned.xlsx")

    print(f"Loading data from {src_file}...")
    df = pd.read_excel(src_file, sheet_name=0)

    # 1. Data Cleaning & Standardization
    # Strip whitespace from string columns
    str_cols = df.select_dtypes(include=["object"]).columns
    for col in str_cols:
        df[col] = df[col].astype(str).str.strip()

    # Standardize DATE format
    df["DATE"] = pd.to_datetime(df["DATE"])
    
    # Sort chronologically by DATE and ORDER_ID
    df = df.sort_values(by=["DATE", "ORDER_ID"]).reset_index(drop=True)

    # Add calculated fields
    df["NET_REVENUE"] = df["REVENUE"] - df["DISCOUNT"]
    df["PROFIT_MARGIN"] = df["PROFIT"] / df["REVENUE"]

    # Reorder columns logically
    cols = [
        "DATE", "ORDER_ID", "CUSTOMER_ID", "REGION", "CITY", "CHANNEL",
        "PRODUCT_CATEGORY", "PRODUCT", "QUANTITY", "UNIT_PRICE", "REVENUE",
        "DISCOUNT", "NET_REVENUE", "MARKETING_COST", "COGS", "PROFIT",
        "PROFIT_MARGIN", "PAYMENT_METHOD", "DELIVERY_TIME_DAYS", "RETURN_FLAG",
        "CUSTOMER_SEGMENT", "SALES_REP"
    ]
    df = df[cols]

    # Convert DATE to formatted string for Excel export
    df_clean_export = df.copy()
    df_clean_export["DATE"] = df_clean_export["DATE"].dt.strftime("%Y-%m-%d")

    # 2. Aggregations for Summary Tables
    # Daily summary
    daily = df.groupby(df["DATE"].dt.strftime("%Y-%m-%d")).agg(
        REVENUE=("REVENUE", "sum"),
        PROFIT=("PROFIT", "sum"),
        NET_REVENUE=("NET_REVENUE", "sum"),
        ORDERS=("ORDER_ID", "count"),
        RETURNS=("RETURN_FLAG", "sum")
    ).reset_index()
    daily.rename(columns={"DATE": "Date"}, inplace=True)
    daily["Margin"] = daily["PROFIT"] / daily["REVENUE"]
    daily["Return_Rate"] = daily["RETURNS"] / daily["ORDERS"]

    # Top 10 Products by Revenue
    top_prod = df.groupby(["PRODUCT", "PRODUCT_CATEGORY"]).agg(
        QUANTITY=("QUANTITY", "sum"),
        REVENUE=("REVENUE", "sum"),
        PROFIT=("PROFIT", "sum"),
        ORDERS=("ORDER_ID", "count")
    ).reset_index().sort_values(by="REVENUE", ascending=False).reset_index(drop=True)
    top_prod["Margin"] = top_prod["PROFIT"] / top_prod["REVENUE"]
    top_10 = top_prod.head(10).copy()
    top_10["Rank"] = range(1, len(top_10) + 1)
    top_10 = top_10[["Rank", "PRODUCT", "PRODUCT_CATEGORY", "QUANTITY", "REVENUE", "PROFIT", "Margin"]]

    # Region x Channel Cross-tabulation
    reg_chn_rev = pd.crosstab(df["REGION"], df["CHANNEL"], values=df["REVENUE"], aggfunc="sum").fillna(0)
    reg_chn_rev["Total_Revenue"] = reg_chn_rev.sum(axis=1)
    reg_profit = df.groupby("REGION")["PROFIT"].sum()
    reg_orders = df.groupby("REGION")["ORDER_ID"].count()
    reg_returns = df.groupby("REGION")["RETURN_FLAG"].sum()
    reg_matrix = reg_chn_rev.copy()
    reg_matrix["Total_Profit"] = reg_profit
    reg_matrix["Margin"] = reg_matrix["Total_Profit"] / reg_matrix["Total_Revenue"]
    reg_matrix["Return_Rate"] = reg_returns / reg_orders
    reg_matrix = reg_matrix.reset_index()

    # Product Category Summary
    cat_summary = df.groupby("PRODUCT_CATEGORY").agg(
        QUANTITY=("QUANTITY", "sum"),
        REVENUE=("REVENUE", "sum"),
        NET_REVENUE=("NET_REVENUE", "sum"),
        PROFIT=("PROFIT", "sum"),
        ORDERS=("ORDER_ID", "count"),
        RETURNS=("RETURN_FLAG", "sum")
    ).reset_index().sort_values(by="REVENUE", ascending=False).reset_index(drop=True)
    cat_summary["Margin"] = cat_summary["PROFIT"] / cat_summary["REVENUE"]
    cat_summary["Return_Rate"] = cat_summary["RETURNS"] / cat_summary["ORDERS"]

    # 3. Create Workbook & Sheets
    wb = openpyxl.Workbook()
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_summary = wb.create_sheet(title="Summary_Tables")
    ws_data = wb.create_sheet(title="Data_Cleaned")

    # Styling definitions
    font_title = Font(name="Segoe UI", size=18, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="D0E0FF")
    font_section = Font(name="Segoe UI", size=12, bold=True, color="1B365D")
    font_kpi_title = Font(name="Segoe UI", size=9, bold=True, color="4A607A")
    font_kpi_value = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
    font_kpi_sub = Font(name="Segoe UI", size=8, color="555555")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=9.5, color="222222")
    font_bold = Font(name="Segoe UI", size=9.5, bold=True, color="1B365D")
    font_callout_h = Font(name="Segoe UI", size=10, bold=True, color="1B365D")
    font_callout_b = Font(name="Segoe UI", size=9, color="333333")

    fill_banner = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_header = PatternFill(start_color="204060", end_color="204060", fill_type="solid")
    fill_sub_header = PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
    fill_kpi = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_kpi_accent = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
    fill_callout = PatternFill(start_color="F9FBFD", end_color="F9FBFD", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")

    thin_border_side = Side(style="thin", color="D3D9E2")
    border_grid = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_kpi = Border(
        left=Side(style="thin", color="BDC9D7"),
        right=Side(style="thin", color="BDC9D7"),
        top=Side(style="medium", color="1B365D"),
        bottom=Side(style="thin", color="BDC9D7")
    )
    border_callout = Border(
        left=Side(style="medium", color="204060"),
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================================
    # SHEET 1: DASHBOARD
    # =========================================================================
    ws_dash.views.sheetView[0].showGridLines = True

    # Banner Header (Rows 1-2)
    ws_dash.merge_cells("B1:M1")
    ws_dash["B1"] = "BÁO CÁO KINH DOANH & DASHBOARD PHÂN TÍCH BÁN HÀNG"
    ws_dash["B1"].font = font_title
    ws_dash["B1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_dash.merge_cells("B2:M2")
    ws_dash["B2"] = "Dữ liệu thực nghiệm 500 đơn hàng (Tháng 06/2024) | Phân tích bởi Business Analyst theo mô hình PDCA"
    ws_dash["B2"].font = font_subtitle
    ws_dash["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for r in range(1, 3):
        for c in range(2, 14):
            ws_dash.cell(row=r, column=c).fill = fill_banner

    # KPI Summary Cards (Rows 4-6)
    kpis = [
        ("TỔNG DOANH THU GỘP", 2827300, "#,##0 ₫", "500 giao dịch (100%)", 2, 3),
        ("DOANH THU THUẦN", 2544994.80, "#,##0 ₫", "Chiết khấu 282,305 (10.0%)", 4, 5),
        ("LỢI NHUẬN RÒNG", 718896.06, "#,##0 ₫", "Biên LNTB: 25.4% | COGS 54.7%", 6, 7),
        ("TỔNG ĐƠN & AOV", 500, "#,##0", "AOV: 5,655 ₫ / đơn", 8, 9),
        ("TỶ LỆ HOÀN HÀNG", 0.256, "0.0%", "128 đơn hoàn (Online 31.4%)", 10, 11),
        ("ĐƠN BỊ LỖ VỐN", 1, "#,##0", "Order ORD00067 (-261 ₫)", 12, 13)
    ]

    for title, val, num_fmt, sub, c_start, c_end in kpis:
        # Title row 4
        ws_dash.merge_cells(start_row=4, start_column=c_start, end_row=4, end_column=c_end)
        cell_t = ws_dash.cell(row=4, column=c_start, value=title)
        cell_t.font = font_kpi_title
        cell_t.alignment = align_center

        # Value row 5
        ws_dash.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_end)
        cell_v = ws_dash.cell(row=5, column=c_start, value=val)
        cell_v.font = font_kpi_value
        cell_v.alignment = align_center
        cell_v.number_format = num_fmt

        # Subtext row 6
        ws_dash.merge_cells(start_row=6, start_column=c_start, end_row=6, end_column=c_end)
        cell_s = ws_dash.cell(row=6, column=c_start, value=sub)
        cell_s.font = font_kpi_sub
        cell_s.alignment = align_center

        # Apply borders and fill
        for r in range(4, 7):
            for c in range(c_start, c_end + 1):
                cell = ws_dash.cell(row=r, column=c)
                cell.fill = fill_kpi_accent if r == 5 else fill_kpi
                cell.border = border_kpi

    # Section 1 Header
    ws_dash.cell(row=8, column=2, value="1. XU HƯỚNG DOANH THU & LỢI NHUẬN THEO THỜI GIAN (06/2024)").font = font_section

    # Chart 1: Line Chart (Doanh thu & Lợi nhuận theo ngày)
    line_chart = LineChart()
    line_chart.title = "Doanh thu và Lợi nhuận qua 30 ngày (Tháng 06/2024)"
    line_chart.style = 13
    line_chart.height = 10
    line_chart.width = 17.5
    line_chart.y_axis.title = "Giá trị (₫)"
    line_chart.x_axis.title = "Ngày trong tháng"

    # Data from Summary_Tables (Cols B & C: Revenue, Profit)
    data_line = Reference(ws_summary, min_col=2, min_row=4, max_col=3, max_row=34)
    cats_line = Reference(ws_summary, min_col=1, min_row=5, max_row=34)
    line_chart.add_data(data_line, titles_from_data=True)
    line_chart.set_categories(cats_line)

    # Style series lines
    if len(line_chart.series) >= 2:
        s1 = line_chart.series[0]
        s1.graphicalProperties.line.solidFill = "1B365D"
        s1.graphicalProperties.line.width = 25000
        s2 = line_chart.series[1]
        s2.graphicalProperties.line.solidFill = "2E8540"
        s2.graphicalProperties.line.width = 25000

    ws_dash.add_chart(line_chart, "B9")

    # Time Insights Box beside Chart 1 (Cols K:M, Rows 9-23)
    ws_dash.merge_cells("K9:M9")
    ws_dash["K9"] = "📌 ĐIỂM NHẤN THỜI GIAN & CHU KỲ"
    ws_dash["K9"].font = font_callout_h
    ws_dash["K9"].alignment = align_left
    ws_dash["K9"].fill = fill_kpi_accent

    time_notes = [
        ("Đỉnh doanh thu (Peak):", "09/06 đạt 153,140 ₫ (20 đơn)"),
        ("Đáy doanh thu (Trough):", "18/06 chỉ đạt 20,840 ₫ (8 đơn)"),
        ("Hiệu ứng Cuối tuần:", "T7 & CN mang lại 1,091,970 ₫"),
        ("Tỷ trọng Cuối tuần:", "Chiếm 38.6% tổng doanh thu"),
        ("AOV Cuối tuần:", "6,500 ₫ (cao hơn 25% ngày thường)"),
        ("Độ ổn định biên LN:", "Biên LN duy trì đều quanh 24-26%"),
        ("Khuyến nghị vận hành:", "Dồn ngân sách Ads & Flashsale vào T6-CN")
    ]
    for idx, (label, val_str) in enumerate(time_notes):
        curr_r = 10 + idx * 2
        ws_dash.merge_cells(start_row=curr_r, start_column=11, end_row=curr_r, end_column=13)
        cell_lbl = ws_dash.cell(row=curr_r, column=11, value=f"{label} {val_str}")
        cell_lbl.font = font_callout_b
        cell_lbl.alignment = align_left
        
        ws_dash.merge_cells(start_row=curr_r+1, start_column=11, end_row=curr_r+1, end_column=13)
        ws_dash.cell(row=curr_r+1, column=11, value="").fill = fill_callout

    for r in range(9, 24):
        for c in range(11, 14):
            ws_dash.cell(row=r, column=c).border = border_grid
            if r == 9:
                ws_dash.cell(row=r, column=c).fill = fill_kpi_accent
            elif r % 2 == 0:
                ws_dash.cell(row=r, column=c).fill = PatternFill(start_color="FFFFFF", fill_type="solid")
            else:
                ws_dash.cell(row=r, column=c).fill = fill_callout

    # Section 2 Header
    ws_dash.cell(row=25, column=2, value="2. SO SÁNH DOANH THU: TOP SẢN PHẨM & KHU VỰC / KÊNH PHÂN PHỐI").font = font_section

    # Chart 2: Top Products Bar Chart (Cols B:G, Rows 27-42)
    bar_prod = BarChart()
    bar_prod.type = "bar"
    bar_prod.title = "Top 10 Sản Phẩm Doanh Thu Cao Nhất (₫)"
    bar_prod.style = 10
    bar_prod.height = 10
    bar_prod.width = 16.5
    bar_prod.y_axis.title = "Sản phẩm"
    bar_prod.x_axis.title = "Doanh thu (₫)"
    bar_prod.legend = None

    data_prod = Reference(ws_summary, min_col=13, min_row=4, max_col=13, max_row=14)
    cats_prod = Reference(ws_summary, min_col=10, min_row=5, max_row=14)
    bar_prod.add_data(data_prod, titles_from_data=True)
    bar_prod.set_categories(cats_prod)
    if bar_prod.series:
        bar_prod.series[0].graphicalProperties.solidFill = "3B5998"

    ws_dash.add_chart(bar_prod, "B27")

    # Chart 3: Clustered Column Chart (Region x Channel) (Cols H:M, Rows 27-42)
    col_reg = BarChart()
    col_reg.type = "col"
    col_reg.grouping = "clustered"
    col_reg.title = "Doanh Thu Theo Khu Vực & Kênh Phân Phối (₫)"
    col_reg.style = 11
    col_reg.height = 10
    col_reg.width = 16.5
    col_reg.y_axis.title = "Doanh thu (₫)"
    col_reg.x_axis.title = "Khu vực"

    # Data from Summary_Tables: Region matrix columns Distributor(R=18), Offline(S=19), Online(T=20)
    data_reg = Reference(ws_summary, min_col=18, min_row=4, max_col=20, max_row=7)
    cats_reg = Reference(ws_summary, min_col=17, min_row=5, max_row=7)
    col_reg.add_data(data_reg, titles_from_data=True)
    col_reg.set_categories(cats_reg)

    if len(col_reg.series) >= 3:
        col_reg.series[0].graphicalProperties.solidFill = "8EA9DB"  # Distributor
        col_reg.series[1].graphicalProperties.solidFill = "1B365D"  # Offline
        col_reg.series[2].graphicalProperties.solidFill = "ED7D31"  # Online

    ws_dash.add_chart(col_reg, "H27")

    # Section 3 Header: Key Strategic Insights & Recommendations (Row 45)
    ws_dash.cell(row=45, column=2, value="3. TỔNG HỢP INSIGHTS TRỌNG TÂM & ĐỀ XUẤT CHIẾN LƯỢC (EXECUTIVE SUMMARY)").font = font_section

    # Insights Card (Cols B:G, Rows 46-56)
    ws_dash.merge_cells("B46:G46")
    ws_dash["B46"] = "🔍 3 BUSINESS INSIGHTS QUAN TRỌNG NHẤT"
    ws_dash["B46"].font = font_callout_h
    ws_dash["B46"].fill = fill_kpi_accent
    ws_dash["B46"].alignment = align_left

    insights_content = [
        "1. Hiệu ứng mua sắm cuối tuần (Weekend Lift): Thứ 7 & CN đóng góp 38.6% doanh thu toàn hệ thống với giá trị đơn hàng vượt trội (+25% AOV).",
        "2. Nghịch lý kênh Online (High Returns & Margin Leakage): Hoàn hàng 31.4% (cao nhất hệ thống), đơn ORD00067 bị lỗ âm 261 ₫ do lạm dụng chiết khấu 19.8% kết hợp phí tiếp thị 15%.",
        "3. Động lực từ Sữa & Đồ uống (High Margin Champions): Sữa (Milk - 28.9% margin) và Cà phê (Coffee - 28.7% margin) là cỗ máy sinh lời lớn nhất, trong khi Cleaner, Cheese chỉ đạt ~20%."
    ]
    for i, ins in enumerate(insights_content):
        r_start = 47 + i * 3
        ws_dash.merge_cells(start_row=r_start, start_column=2, end_row=r_start+2, end_column=7)
        cell = ws_dash.cell(row=r_start, column=2, value=ins)
        cell.font = font_callout_b
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for r in range(46, 56):
        for c in range(2, 8):
            ws_dash.cell(row=r, column=c).border = border_callout
            if r == 46:
                ws_dash.cell(row=r, column=c).fill = fill_kpi_accent
            else:
                ws_dash.cell(row=r, column=c).fill = fill_callout

    # Actions Card (Cols H:M, Rows 46-56)
    ws_dash.merge_cells("H46:M46")
    ws_dash["H46"] = "🚀 2 ĐỀ XUẤT HÀNH ĐỘNG CHIẾN LƯỢC"
    ws_dash["H46"].font = font_callout_h
    ws_dash["H46"].fill = PatternFill(start_color="E2F0D9", fill_type="solid")
    ws_dash["H46"].alignment = align_left

    actions_content = [
        "HÀNH ĐỘNG 1: Kiểm soát trần chiết khấu & Tối ưu tỷ lệ hoàn hàng Online\n• Đặt trần chiết khấu Online tối đa 12% để bảo vệ Gross Margin > 20%.\n• Kiểm toán nguyên nhân hoàn hàng 31.4%, nâng cấp đóng gói để giảm hoàn về mức < 20%.",
        "HÀNH ĐỘNG 2: Tung chiến dịch 'Weekend Super-Booster' & Cross-Selling Bundle\n• Tối ưu lịch quảng cáo tập trung từ chiều Thứ 6 đến Chủ Nhật.\n• Đóng gói combo sản phẩm bán chạy (Juice, Milk) đi kèm sản phẩm tiêu thụ chậm (Cheese, Cleaner) nhằm đẩy nhanh hàng tồn kho."
    ]
    for i, act in enumerate(actions_content):
        r_start = 47 + i * 4 + (1 if i == 1 else 0)
        r_end = r_start + 3
        ws_dash.merge_cells(start_row=r_start, start_column=8, end_row=r_end, end_column=13)
        cell = ws_dash.cell(row=r_start, column=8, value=act)
        cell.font = font_callout_b
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for r in range(46, 56):
        for c in range(8, 14):
            ws_dash.cell(row=r, column=c).border = border_grid
            if r == 46:
                ws_dash.cell(row=r, column=c).fill = PatternFill(start_color="E2F0D9", fill_type="solid")
            else:
                ws_dash.cell(row=r, column=c).fill = PatternFill(start_color="F9FCF7", fill_type="solid")

    # =========================================================================
    # SHEET 2: SUMMARY_TABLES
    # =========================================================================
    ws_summary.views.sheetView[0].showGridLines = True

    # Title
    ws_summary.merge_cells("A1:U1")
    ws_summary["A1"] = "BẢNG TỔNG HỢP NGUỒN DỮ LIỆU DASHBOARD & BÁO CÁO PIVOT"
    ws_summary["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    ws_summary["A1"].fill = fill_banner
    ws_summary["A1"].alignment = align_left

    # --- TABLE 1: Daily Trend (Cols A:G) ---
    ws_summary.merge_cells("A3:G3")
    ws_summary["A3"] = "BẢNG 1: DOANH THU & LỢI NHUẬN THEO NGÀY (THÁNG 06/2024)"
    ws_summary["A3"].font = font_section

    daily_headers = ["Ngày", "Doanh thu (₫)", "Lợi nhuận (₫)", "Doanh thu thuần (₫)", "Số đơn", "Hoàn hàng", "Biên LN (%)"]
    for c_idx, h in enumerate(daily_headers, 1):
        cell = ws_summary.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_grid

    for r_idx, row_data in daily.iterrows():
        r = 5 + r_idx
        ws_summary.cell(row=r, column=1, value=row_data["Date"]).alignment = align_center
        ws_summary.cell(row=r, column=2, value=row_data["REVENUE"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=3, value=row_data["PROFIT"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=4, value=row_data["NET_REVENUE"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=5, value=row_data["ORDERS"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=6, value=row_data["RETURNS"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=7, value=row_data["Margin"]).number_format = "0.0%"

        for c_idx in range(1, 8):
            cell = ws_summary.cell(row=r, column=c_idx)
            cell.font = font_body
            cell.border = border_grid
            if r % 2 == 0:
                cell.fill = fill_zebra

    # Total Row Daily
    tot_r_daily = 5 + len(daily)
    ws_summary.cell(row=tot_r_daily, column=1, value="Tổng cộng").font = font_bold
    ws_summary.cell(row=tot_r_daily, column=1).alignment = align_center
    ws_summary.cell(row=tot_r_daily, column=2, value=f"=SUM(B5:B{tot_r_daily-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_daily, column=3, value=f"=SUM(C5:C{tot_r_daily-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_daily, column=4, value=f"=SUM(D5:D{tot_r_daily-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_daily, column=5, value=f"=SUM(E5:E{tot_r_daily-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_daily, column=6, value=f"=SUM(F5:F{tot_r_daily-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_daily, column=7, value=f"=C{tot_r_daily}/B{tot_r_daily}").number_format = "0.0%"

    for c_idx in range(1, 8):
        cell = ws_summary.cell(row=tot_r_daily, column=c_idx)
        cell.font = font_bold
        cell.border = border_grid
        cell.fill = fill_kpi_accent

    # --- TABLE 2: Top 10 Products (Cols I:O) ---
    ws_summary.merge_cells("I3:O3")
    ws_summary["I3"] = "BẢNG 2: TOP 10 SẢN PHẨM DOANH THU CAO NHẤT"
    ws_summary["I3"].font = font_section

    prod_headers = ["Hạng", "Sản phẩm", "Ngành hàng", "Số lượng", "Doanh thu (₫)", "Lợi nhuận (₫)", "Biên LN (%)"]
    for c_idx, h in enumerate(prod_headers, 9):
        cell = ws_summary.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_grid

    for r_idx, row_data in top_10.iterrows():
        r = 5 + r_idx
        ws_summary.cell(row=r, column=9, value=row_data["Rank"]).alignment = align_center
        ws_summary.cell(row=r, column=10, value=row_data["PRODUCT"]).alignment = align_left
        ws_summary.cell(row=r, column=11, value=row_data["PRODUCT_CATEGORY"]).alignment = align_left
        ws_summary.cell(row=r, column=12, value=row_data["QUANTITY"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=13, value=row_data["REVENUE"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=14, value=row_data["PROFIT"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=15, value=row_data["Margin"]).number_format = "0.0%"

        for c_idx in range(9, 16):
            cell = ws_summary.cell(row=r, column=c_idx)
            cell.font = font_body
            cell.border = border_grid
            if r % 2 == 0:
                cell.fill = fill_zebra

    # Total Row Top 10
    tot_r_top = 5 + len(top_10)
    ws_summary.cell(row=tot_r_top, column=10, value="Tổng Top 10").font = font_bold
    ws_summary.cell(row=tot_r_top, column=12, value=f"=SUM(L5:L{tot_r_top-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_top, column=13, value=f"=SUM(M5:M{tot_r_top-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_top, column=14, value=f"=SUM(N5:N{tot_r_top-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_top, column=15, value=f"=N{tot_r_top}/M{tot_r_top}").number_format = "0.0%"

    for c_idx in range(9, 16):
        cell = ws_summary.cell(row=tot_r_top, column=c_idx)
        cell.font = font_bold
        cell.border = border_grid
        cell.fill = fill_kpi_accent

    # --- TABLE 3: Region x Channel Breakdown (Cols Q:W) ---
    ws_summary.merge_cells("Q3:W3")
    ws_summary["Q3"] = "BẢNG 3: MA TRẬN DOANH THU THEO KHU VỰC & KÊNH BÁN HÀNG"
    ws_summary["Q3"].font = font_section

    reg_headers = ["Khu vực", "Distributor (₫)", "Offline (₫)", "Online (₫)", "Tổng Doanh Thu", "Tổng Lợi Nhuận", "Biên LN (%)"]
    for c_idx, h in enumerate(reg_headers, 17):
        cell = ws_summary.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_grid

    for r_idx, row_data in reg_matrix.iterrows():
        r = 5 + r_idx
        ws_summary.cell(row=r, column=17, value=row_data["REGION"]).alignment = align_center
        ws_summary.cell(row=r, column=18, value=row_data["Distributor"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=19, value=row_data["Offline"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=20, value=row_data["Online"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=21, value=row_data["Total_Revenue"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=22, value=row_data["Total_Profit"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=23, value=row_data["Margin"]).number_format = "0.0%"

        for c_idx in range(17, 24):
            cell = ws_summary.cell(row=r, column=c_idx)
            cell.font = font_body
            cell.border = border_grid

    tot_r_reg = 5 + len(reg_matrix)
    ws_summary.cell(row=tot_r_reg, column=17, value="Toàn quốc").font = font_bold
    ws_summary.cell(row=tot_r_reg, column=17).alignment = align_center
    ws_summary.cell(row=tot_r_reg, column=18, value=f"=SUM(R5:R{tot_r_reg-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_reg, column=19, value=f"=SUM(S5:S{tot_r_reg-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_reg, column=20, value=f"=SUM(T5:T{tot_r_reg-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_reg, column=21, value=f"=SUM(U5:U{tot_r_reg-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_reg, column=22, value=f"=SUM(V5:V{tot_r_reg-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_reg, column=23, value=f"=V{tot_r_reg}/U{tot_r_reg}").number_format = "0.0%"

    for c_idx in range(17, 24):
        cell = ws_summary.cell(row=tot_r_reg, column=c_idx)
        cell.font = font_bold
        cell.border = border_grid
        cell.fill = fill_kpi_accent

    # --- TABLE 4: Product Category Performance (Cols I:P, Row 18) ---
    ws_summary.merge_cells("I18:P18")
    ws_summary["I18"] = "BẢNG 4: HIỆU SUẤT THEO NGÀNH HÀNG (PRODUCT CATEGORIES)"
    ws_summary["I18"].font = font_section

    cat_headers = ["Ngành hàng", "Số lượng", "Doanh thu (₫)", "Doanh thu thuần", "Lợi nhuận (₫)", "Biên LN (%)", "Số đơn", "Tỷ lệ hoàn"]
    for c_idx, h in enumerate(cat_headers, 9):
        cell = ws_summary.cell(row=19, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = border_grid

    for r_idx, row_data in cat_summary.iterrows():
        r = 20 + r_idx
        ws_summary.cell(row=r, column=9, value=row_data["PRODUCT_CATEGORY"]).alignment = align_left
        ws_summary.cell(row=r, column=10, value=row_data["QUANTITY"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=11, value=row_data["REVENUE"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=12, value=row_data["NET_REVENUE"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=13, value=row_data["PROFIT"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=14, value=row_data["Margin"]).number_format = "0.0%"
        ws_summary.cell(row=r, column=15, value=row_data["ORDERS"]).number_format = "#,##0"
        ws_summary.cell(row=r, column=16, value=row_data["Return_Rate"]).number_format = "0.0%"

        for c_idx in range(9, 17):
            cell = ws_summary.cell(row=r, column=c_idx)
            cell.font = font_body
            cell.border = border_grid
            if r % 2 == 0:
                cell.fill = fill_zebra

    tot_r_cat = 20 + len(cat_summary)
    ws_summary.cell(row=tot_r_cat, column=9, value="Tổng toàn bộ").font = font_bold
    ws_summary.cell(row=tot_r_cat, column=10, value=f"=SUM(J20:J{tot_r_cat-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_cat, column=11, value=f"=SUM(K20:K{tot_r_cat-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_cat, column=12, value=f"=SUM(L20:L{tot_r_cat-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_cat, column=13, value=f"=SUM(M20:M{tot_r_cat-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_cat, column=14, value=f"=M{tot_r_cat}/K{tot_r_cat}").number_format = "0.0%"
    ws_summary.cell(row=tot_r_cat, column=15, value=f"=SUM(O20:O{tot_r_cat-1})").number_format = "#,##0"
    ws_summary.cell(row=tot_r_cat, column=16, value=f"=AVERAGE(P20:P{tot_r_cat-1})").number_format = "0.0%"

    for c_idx in range(9, 17):
        cell = ws_summary.cell(row=tot_r_cat, column=c_idx)
        cell.font = font_bold
        cell.border = border_grid
        cell.fill = fill_kpi_accent

    # Auto-adjust column widths for Summary_Tables
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # =========================================================================
    # SHEET 3: DATA_CLEANED
    # =========================================================================
    ws_data.views.sheetView[0].showGridLines = True

    # Headers
    data_headers = df_clean_export.columns.tolist()
    for c_idx, h in enumerate(data_headers, 1):
        cell = ws_data.cell(row=1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_grid

    # Rows
    for r_idx, row_series in df_clean_export.iterrows():
        r = 2 + r_idx
        for c_idx, val in enumerate(row_series, 1):
            cell = ws_data.cell(row=r, column=c_idx, value=val)
            cell.font = font_body
            cell.border = border_grid

            # Number formats
            col_name = data_headers[c_idx - 1]
            if col_name in ["UNIT_PRICE", "REVENUE", "DISCOUNT", "NET_REVENUE", "MARKETING_COST", "COGS", "PROFIT"]:
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif col_name == "PROFIT_MARGIN":
                cell.number_format = "0.0%"
                cell.alignment = align_right
            elif col_name in ["QUANTITY", "DELIVERY_TIME_DAYS", "RETURN_FLAG"]:
                cell.number_format = "#,##0"
                cell.alignment = align_center
            elif col_name == "DATE":
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if r % 2 == 1:
                cell.fill = fill_zebra

    # Freeze Header & Add AutoFilter
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = f"A1:V{len(df_clean_export) + 1}"

    # Auto-adjust column widths for Data_Cleaned
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or "")) for cell in col[:15])
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Adjust Dashboard column widths
    ws_dash.column_dimensions["A"].width = 3
    ws_dash.column_dimensions["B"].width = 15
    ws_dash.column_dimensions["C"].width = 15
    ws_dash.column_dimensions["D"].width = 15
    ws_dash.column_dimensions["E"].width = 15
    ws_dash.column_dimensions["F"].width = 15
    ws_dash.column_dimensions["G"].width = 15
    ws_dash.column_dimensions["H"].width = 15
    ws_dash.column_dimensions["I"].width = 15
    ws_dash.column_dimensions["J"].width = 15
    ws_dash.column_dimensions["K"].width = 15
    ws_dash.column_dimensions["L"].width = 15
    ws_dash.column_dimensions["M"].width = 15

    print(f"Saving workbook to {out_file}...")
    wb.save(out_file)
    print("Workbook successfully created!")

if __name__ == "__main__":
    generate_sales_dashboard()
